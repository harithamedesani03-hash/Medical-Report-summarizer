"""
Document processing utilities: extract text from PDF, DOCX, TXT, and images.
"""

import io
import pymupdf
import docx
from PIL import Image


def extract_text_from_pdf(file_bytes: bytes) -> tuple[str, list]:
    """
    Extract text and page images from a PDF.
    Returns (full_text, list_of_PIL_images).
    """
    text_parts = []
    images = []

    pdf_doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    for page_num in range(len(pdf_doc)):
        page = pdf_doc[page_num]

        # Extract text
        page_text = page.get_text("text")
        if page_text.strip():
            text_parts.append(f"[Page {page_num + 1}]\n{page_text}")

        # Render page as image (for vision fallback)
        mat = pymupdf.Matrix(2.0, 2.0)  # 2x zoom for better resolution
        pix = page.get_pixmap(matrix=mat)
        img_bytes = pix.tobytes("png")
        images.append(Image.open(io.BytesIO(img_bytes)))

    pdf_doc.close()
    full_text = "\n\n".join(text_parts)
    return full_text, images


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a .docx file."""
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also extract tables
    table_texts = []
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                table_texts.append(row_text)

    all_text = "\n".join(paragraphs)
    if table_texts:
        all_text += "\n\n[Tables]\n" + "\n".join(table_texts)

    return all_text


def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extract text from a plain text file."""
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return file_bytes.decode("latin-1", errors="replace")


def load_image(file_bytes: bytes) -> Image.Image:
    """Load image from bytes."""
    return Image.open(io.BytesIO(file_bytes))


def process_uploaded_file(file_bytes: bytes, file_name: str) -> tuple[str, list, str]:
    """
    Process any supported file type.
    Returns (extracted_text, list_of_images, file_type_label).
    
    - PDF  → text extraction + page images
    - DOCX → text extraction
    - TXT  → raw text
    - Image → empty text + single image
    """
    name_lower = file_name.lower()

    if name_lower.endswith(".pdf"):
        text, images = extract_text_from_pdf(file_bytes)
        return text, images, "PDF"

    elif name_lower.endswith(".docx"):
        text = extract_text_from_docx(file_bytes)
        return text, [], "DOCX"

    elif name_lower.endswith(".txt"):
        text = extract_text_from_txt(file_bytes)
        return text, [], "TXT"

    elif name_lower.endswith((".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".webp")):
        image = load_image(file_bytes)
        return "", [image], "Image"

    else:
        raise ValueError(f"Unsupported file type: {file_name}")
